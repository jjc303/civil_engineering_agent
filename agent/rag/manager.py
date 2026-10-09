from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any
from uuid import uuid4

from agent.contracts.knowledge import DocumentType
from .chroma_adapter import ChromaKnowledgeRetriever, KnowledgeChunk


@dataclass(frozen=True)
class PreparedDocument:
    records: list[dict[str, Any]]
    page_count: int | None
    parser_name: str


class RagManager:
    """Shared parsing, chunking, embedding and vector access for document types."""

    def __init__(self, retriever: ChromaKnowledgeRetriever, chunk_size: int, chunk_overlap: int) -> None:
        self.retriever = retriever
        self.chunk_size = chunk_size
        self.chunk_overlap = chunk_overlap

    def prepare(self, path: Path, document_type: DocumentType, metadata: dict[str, Any]) -> PreparedDocument:
        text, page_count, parser = _extract_text(path)
        if document_type == "STANDARD":
            chunks = _standard_chunks(path, self.chunk_size, self.chunk_overlap)
        else:
            chunks = [(body, f"段落 {index + 1}") for index, body in enumerate(_chunks(text, self.chunk_size, self.chunk_overlap))]
        if not chunks:
            raise ValueError("document contains no indexable text")
        records = [
            {
                "chunk_id": str(uuid4()),
                "text": body,
                "metadata": {**metadata, "document_type": document_type, "page_or_section": section, "chunk_index": index, "status": "ACTIVE"},
            }
            for index, (body, section) in enumerate(chunks)
        ]
        return PreparedDocument(records, page_count, parser)

    def replace_version(self, document_type: DocumentType, version_id: str, records: list[dict[str, Any]]) -> None:
        if document_type == "STANDARD":
            self.retriever.replace_standard_version(version_id, records)
        else:
            self.retriever.replace_version(version_id, records)

    def delete_version(self, document_type: DocumentType, version_id: str) -> None:
        if document_type == "STANDARD":
            self.retriever.delete_standard_version(version_id)
        else:
            self.retriever.delete_version(version_id)

    def count(self, document_type: DocumentType) -> int:
        if document_type == "STANDARD":
            return self.retriever.count_standards() if hasattr(self.retriever, "count_standards") else 0
        return self.retriever.count()

    def search(self, document_type: DocumentType, query: str, top_k: int, version_ids: list[str] | None = None) -> list[KnowledgeChunk]:
        if version_ids is not None and hasattr(self.retriever, "search_active"):
            return self.retriever.search_active(document_type, query, top_k, version_ids)
        if document_type == "STANDARD":
            return self.retriever.search_standards(query, top_k)
        return self.retriever.search(query, top_k)

    def get_version(self, document_type: DocumentType, version_id: str) -> list[KnowledgeChunk]:
        return self.retriever.get_version(document_type, version_id)

    def update_version_source(self, document_type: DocumentType, version_id: str, source_label: str) -> None:
        if hasattr(self.retriever, "update_version_source"):
            self.retriever.update_version_source(document_type, version_id, source_label)


def _extract_text(path: Path) -> tuple[str, int | None, str]:
    suffix = path.suffix.lower()
    if suffix in {".txt", ".md"}:
        return path.read_text(encoding="utf-8"), None, "plain_text"
    if suffix == ".pdf":
        from pypdf import PdfReader
        pages = PdfReader(str(path)).pages
        return "\n".join(page.extract_text() or "" for page in pages), len(pages), "pypdf"
    if suffix == ".docx":
        from docx import Document
        return "\n".join(p.text for p in Document(str(path)).paragraphs), None, "python-docx"
    raise ValueError("unsupported knowledge document")


def _chunks(text: str, size: int, overlap: int) -> list[str]:
    text, out, start = " ".join(text.split()), [], 0
    while start < len(text):
        end = min(len(text), start + size)
        out.append(text[start:end])
        if end == len(text):
            break
        start = end - overlap
    return out


_SECTION_HEADING = re.compile(r"^(?:#{1,6}\s+|第[一二三四五六七八九十百千\d]+[章节条]\s*|\d+(?:\.\d+)*[.、]?\s+).{1,100}$")


def _standard_chunks(path: Path, size: int, overlap: int) -> list[tuple[str, str]]:
    """Keep section headings and PDF page locations attached to vector chunks."""
    if path.suffix.lower() == ".pdf":
        from pypdf import PdfReader
        sources = [(page.extract_text() or "", f"第 {index} 页") for index, page in enumerate(PdfReader(str(path)).pages, 1)]
    elif path.suffix.lower() == ".docx":
        from docx import Document
        paragraphs = Document(str(path)).paragraphs
        sources = [("\n".join(("# " if p.style and p.style.name.lower().startswith("heading") else "") + p.text for p in paragraphs), "文档")]
    else:
        sources = [(path.read_text(encoding="utf-8"), "文档")]
    output: list[tuple[str, str]] = []
    for source_text, location in sources:
        section = location
        lines: list[str] = []

        def flush() -> None:
            if not lines:
                return
            body = "\n".join(lines).strip()
            for part in _chunks(body, size, overlap):
                if part.strip():
                    output.append((part, section))

        for raw_line in source_text.splitlines():
            line = raw_line.strip()
            if not line:
                continue
            if _SECTION_HEADING.match(line):
                flush()
                section = f"{location} · {line.lstrip('# ').strip()}" if location != "文档" else line.lstrip("# ").strip()
                lines = [line]
            else:
                lines.append(line)
        flush()
    return output
